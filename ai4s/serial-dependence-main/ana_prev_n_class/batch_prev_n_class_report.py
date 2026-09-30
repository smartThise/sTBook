"""Batch prev-n-class analysis + Word report generation.

Repo assumptions
----------------
- Run outputs live under: <project_root>/outputs/<run_tag>/
- Each run contains: run_config.yaml (for metadata)
- The analyzer is: <project_root>/ana_prev_n_class/analyze_prevclass_bar.py
- Analyzer outputs live under:
    <project_root>/ana_prev_n_class/results/<run_tag>/
      - n_back_stats.csv
      - delta_vs_n.png (optional)
      - meta.json

What this script does
---------------------
1) Scan outputs/* run folders.
2) Select eligible runs: dimension NOT in {none, 2d, mixed}.
3) For each eligible run, run the analyzer (if needed):
     python ana_prev_n_class/analyze_prevclass_bar.py --run_dir outputs/<run_tag> --params ana_prev_n_class/params.yaml
4) Generate a Word report (docx) under:
     <project_root>/ana_prev_n_class/<report_name>

Usage
-----
  python ana_prev_n_class/batch_prev_n_class_report.py
  python ana_prev_n_class/batch_prev_n_class_report.py --dry-run
  python ana_prev_n_class/batch_prev_n_class_report.py --force
  python ana_prev_n_class/batch_prev_n_class_report.py --params ana_prev_n_class/params.yaml
"""

from __future__ import annotations

import argparse
import datetime as _dt
import os
import re
import subprocess
import sys
from dataclasses import dataclass
from typing import Any, Dict, List, Optional

import yaml


SKIP_DIMS = ("none", "2d", "mixed")


# --------------------------- helpers ---------------------------

def project_root_dir() -> str:
    """Best-effort project root detection."""
    here = os.path.abspath(os.path.dirname(__file__))  # .../ana_prev_n_class
    parent = os.path.abspath(os.path.join(here, ".."))
    if os.path.isdir(os.path.join(parent, "outputs")) or os.path.exists(os.path.join(parent, "config.yaml")):
        return parent
    if os.path.isdir(os.path.join(here, "outputs")) or os.path.exists(os.path.join(here, "config.yaml")):
        return here
    return parent


def safe_load_yaml(path: str) -> Dict[str, Any]:
    try:
        with open(path, "r", encoding="utf-8") as f:
            obj = yaml.safe_load(f) or {}
        if isinstance(obj, dict):
            return obj
    except Exception:
        pass
    return {}


def parse_n_from_run_tag(run_tag: str) -> Optional[int]:
    m = re.search(r"(?:^|_)n(\d+)(?:_|$)", run_tag)
    if not m:
        return None
    try:
        return int(m.group(1))
    except Exception:
        return None


def infer_model_from_run_tag(run_tag: str) -> Optional[str]:
    parts = run_tag.split("_")
    if len(parts) < 4:
        return None
    third = parts[2]
    if third in ("exp1", "exp2", "both"):
        return None
    return third


def infer_which_from_run_tag(run_tag: str) -> Optional[str]:
    parts = run_tag.split("_")
    if len(parts) < 4:
        return None
    third = parts[2]
    if third in ("exp1", "exp2", "both"):
        return third
    fourth = parts[3]
    if fourth in ("exp1", "exp2", "both"):
        return fourth
    return None


@dataclass
class RunMeta:
    run_tag: str
    run_dir: str
    model: str
    which: str
    dimension: str
    n_words: Optional[int]
    n_words_exp2: Optional[int]
    temperature: Optional[float]
    seed: Optional[int]


def extract_run_meta(run_dir: str) -> Optional[RunMeta]:
    run_tag = os.path.basename(os.path.normpath(run_dir))
    cfg_path = os.path.join(run_dir, "run_config.yaml")
    if not os.path.exists(cfg_path):
        return None

    cfg = safe_load_yaml(cfg_path)
    merged = cfg.get("merged_cfg") if isinstance(cfg.get("merged_cfg"), dict) else {}

    def _pick(key: str, default: Any = None) -> Any:
        v = cfg.get(key, None)
        if v is None and isinstance(merged, dict):
            v = merged.get(key, None)
        return default if v is None else v

    dimension = str(_pick("dimension", "unknown")).strip().lower()
    if dimension in SKIP_DIMS:
        return None

    which = str(_pick("which", "")).strip().lower() or (infer_which_from_run_tag(run_tag) or "unknown")
    model = str(_pick("model", "")).strip() or (infer_model_from_run_tag(run_tag) or "(unknown)")

    n_words = None
    if _pick("n_words", None) is not None:
        try:
            n_words = int(_pick("n_words"))
        except Exception:
            n_words = None
    if n_words is None:
        n_words = parse_n_from_run_tag(run_tag)

    n_words_exp2 = None
    if _pick("n_words_exp2", None) is not None:
        try:
            n_words_exp2 = int(_pick("n_words_exp2"))
        except Exception:
            n_words_exp2 = None

    temperature = None
    if _pick("temperature", None) is not None:
        try:
            temperature = float(_pick("temperature"))
        except Exception:
            temperature = None

    seed = None
    if _pick("seed", None) is not None:
        try:
            seed = int(_pick("seed"))
        except Exception:
            seed = None

    return RunMeta(
        run_tag=run_tag,
        run_dir=run_dir,
        model=model,
        which=which,
        dimension=dimension,
        n_words=n_words,
        n_words_exp2=n_words_exp2,
        temperature=temperature,
        seed=seed,
    )


def has_exp2_data(run_dir: str) -> bool:
    return os.path.exists(os.path.join(run_dir, "exp2_sequential_raw.csv")) or os.path.exists(
        os.path.join(run_dir, "combined_overview.csv")
    )


def analysis_dir(root: str, run_tag: str) -> str:
    return os.path.join(root, "ana_prev_n_class", "results", run_tag)


def analysis_paths(root: str, run_tag: str) -> Dict[str, str]:
    d = analysis_dir(root, run_tag)
    return {
        "stats_csv": os.path.join(d, "n_back_stats.csv"),
        "fig_png": os.path.join(d, "delta_vs_n.png"),
        "meta_json": os.path.join(d, "meta.json"),
    }


def ensure_analysis(root: str, meta: RunMeta, params_path: str, force: bool) -> Optional[str]:
    """Return error message if failed, else None."""
    paths = analysis_paths(root, meta.run_tag)
    if not force and os.path.exists(paths["stats_csv"]) and os.path.exists(paths["meta_json"]):
        # figure may be optional; don't require.
        return None

    analyzer = os.path.join(root, "ana_prev_n_class", "analyze_prevclass_bar.py")
    if not os.path.exists(analyzer):
        return f"Analyzer not found: {analyzer}"

    cmd = [
        sys.executable,
        analyzer,
        "--run_dir",
        meta.run_dir,
        "--params",
        params_path,
    ]

    try:
        subprocess.run(cmd, cwd=root, check=True)
    except subprocess.CalledProcessError as e:
        return f"analyze_prevclass_bar failed (run={meta.run_tag}): {e}"

    if not os.path.exists(paths["stats_csv"]) or not os.path.exists(paths["meta_json"]):
        return f"Analysis did not produce expected outputs for run={meta.run_tag}"

    return None


@dataclass
class AnalysisRecord:
    meta: RunMeta
    out_dir: str
    stats_csv: Optional[str]
    fig_png: Optional[str]
    rows_preview: Optional[List[Dict[str, Any]]] = None
    note: Optional[str] = None


def load_record(root: str, meta: RunMeta) -> AnalysisRecord:
    import pandas as pd

    out_dir = analysis_dir(root, meta.run_tag)
    paths = analysis_paths(root, meta.run_tag)

    stats_csv = paths["stats_csv"] if os.path.exists(paths["stats_csv"]) else None
    fig_png = paths["fig_png"] if os.path.exists(paths["fig_png"]) else None

    rows_preview = None
    note = None

    if stats_csv:
        try:
            df = pd.read_csv(stats_csv)
            cols = [c for c in ["n", "delta_high_minus_low", "p_value", "n_low", "n_high", "n_pairs", "test"] if c in df.columns]
            rows_preview = df[cols].sort_values("n").to_dict(orient="records")
        except Exception as e:
            note = f"Failed to parse stats CSV: {e}"

    return AnalysisRecord(meta=meta, out_dir=out_dir, stats_csv=stats_csv, fig_png=fig_png, rows_preview=rows_preview, note=note)


def build_docx_report(records: List[AnalysisRecord], out_path: str, params_path: str) -> None:
    try:
        from docx import Document  # type: ignore
        from docx.oxml import OxmlElement  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.shared import Inches  # type: ignore
    except Exception as e:
        raise RuntimeError(
            "python-docx is required to generate a Word report. "
            "Install it via: pip install python-docx\n"
            f"Import error: {e}"
        )

    def _set_cell_width(cell, width_in_inches: float) -> None:
        """Force cell width via tcW so Word actually respects it."""
        width = Inches(width_in_inches)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)
        tcW.set(qn("w:w"), str(width.twips))  # twips
        tcW.set(qn("w:type"), "dxa")

    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    doc = Document()
    doc.add_heading("Prev-n-class (low/high) n-back analysis report", level=0)
    doc.add_paragraph(f"Generated at: {_dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # include sweep params at top (helps reproducibility)
    sweep = safe_load_yaml(params_path)
    doc.add_paragraph(f"Params file: {os.path.abspath(params_path)}")
    if sweep:
        lines = []
        if "n_values" in sweep:
            lines.append(f"n_values: {sweep.get('n_values')}")
        lines.append(f"low_max: {sweep.get('low_max', 50)}")
        lines.append(f"high_min: {sweep.get('high_min', 51)}")
        lines.append(f"test: {sweep.get('test', 'welch')}")
        if str(sweep.get("test", "welch")).lower() == "permutation":
            lines.append(f"n_perm: {sweep.get('n_perm', 0)}")
            lines.append(f"seed: {sweep.get('seed', 0)}")
        doc.add_paragraph("\n".join(lines))

    if not records:
        doc.add_paragraph("(No eligible runs found.)")
        doc.save(out_path)
        return

    # stable order
    records = sorted(records, key=lambda r: r.meta.run_tag)

    # pre-compute usable width from section settings (inches)
    sec = doc.sections[0]
    page_w_in = sec.page_width.inches
    margin_l_in = sec.left_margin.inches
    margin_r_in = sec.right_margin.inches
    usable_w_in = max(0.0, page_w_in - margin_l_in - margin_r_in)

    for rec in records:
        m = rec.meta
        doc.add_heading(m.run_tag, level=1)

        # 1x2 table: [image | metadata]
        tbl = doc.add_table(rows=1, cols=2)
        tbl.autofit = False

        left = tbl.cell(0, 0)
        right = tbl.cell(0, 1)

        # Give right column enough room for text; left gets the rest.
        # right_w_in can be tuned; this keeps metadata readable while preventing overlap.
        right_w_in = min(2.35, max(1.9, usable_w_in * 0.33))
        left_w_in = max(1.0, usable_w_in - right_w_in)

        _set_cell_width(left, left_w_in)
        _set_cell_width(right, right_w_in)

        # image (must be <= left column width)
        pic_w_in = max(1.0, left_w_in - 0.25)  # leave margin to avoid overflow due to cell padding
        if rec.fig_png and os.path.exists(rec.fig_png):
            p = left.paragraphs[0]
            run = p.add_run()
            run.add_picture(rec.fig_png, width=Inches(pic_w_in))
        else:
            left.text = "(missing delta_vs_n.png)"

        # metadata
        lines_meta: List[str] = []
        lines_meta.append(f"model: {m.model}")
        lines_meta.append(f"which: {m.which}")
        lines_meta.append(f"dimension: {m.dimension}")
        if m.n_words is not None:
            lines_meta.append(f"n_words: {m.n_words}")
        if m.n_words_exp2 is not None:
            lines_meta.append(f"n_words_exp2: {m.n_words_exp2}")
        if m.temperature is not None:
            lines_meta.append(f"temperature: {m.temperature}")
        if m.seed is not None:
            lines_meta.append(f"seed: {m.seed}")
        if rec.stats_csv:
            lines_meta.append(f"stats_csv: {os.path.relpath(rec.stats_csv, project_root_dir())}")
        if rec.note:
            lines_meta.append(f"note: {rec.note}")

        right.text = "\n".join(lines_meta)

        # stats table
        if rec.rows_preview:
            doc.add_paragraph("n-back deltas (per n):")
            headers = list(rec.rows_preview[0].keys())
            t = doc.add_table(rows=1, cols=len(headers))
            for j, h in enumerate(headers):
                t.cell(0, j).text = str(h)
            for row in rec.rows_preview:
                tr = t.add_row().cells
                for j, h in enumerate(headers):
                    v = row.get(h)
                    if isinstance(v, float):
                        if h == "p_value":
                            tr[j].text = f"{v:.3g}"
                        else:
                            tr[j].text = f"{v:.6f}"
                    else:
                        tr[j].text = "" if v is None else str(v)

        doc.add_paragraph("")

    doc.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser(description="Scan outputs, run prev-n-class analysis, and generate a Word report.")
    ap.add_argument("--outputs-dir", type=str, default="outputs", help="Outputs directory under project root.")
    ap.add_argument(
        "--params",
        type=str,
        default=os.path.join("ana_prev_n_class", "params.yaml"),
        help="Sweep params YAML passed to analyzer.",
    )
    ap.add_argument(
        "--report-name",
        type=str,
        default="prev_n_class_report.docx",
        help="Report filename to write under ana_prev_n_class/.",
    )
    ap.add_argument("--dry-run", action="store_true", help="Only scan and print what would be done.")
    ap.add_argument("--force", action="store_true", help="Force rerun analysis for all eligible runs.")
    args = ap.parse_args()

    root = project_root_dir()
    outputs_dir = os.path.join(root, args.outputs_dir)
    if not os.path.isdir(outputs_dir):
        raise FileNotFoundError(f"Outputs directory not found: {outputs_dir}")

    params_path = args.params
    if not os.path.isabs(params_path):
        params_path = os.path.join(root, params_path)

    metas: List[RunMeta] = []
    for name in sorted(os.listdir(outputs_dir)):
        run_dir = os.path.join(outputs_dir, name)
        if not os.path.isdir(run_dir):
            continue
        meta = extract_run_meta(run_dir)
        if meta is None:
            continue
        if not has_exp2_data(run_dir):
            continue
        metas.append(meta)

    if not metas:
        print("[WARN] No eligible runs found under outputs/.")

    failures: List[str] = []
    ensured: List[RunMeta] = []

    for m in metas:
        if args.dry_run:
            p = analysis_paths(root, m.run_tag)
            already = os.path.exists(p["stats_csv"]) and os.path.exists(p["meta_json"])
            if already and not args.force:
                print(f"[SKIP] already analyzed: {m.run_tag}")
            else:
                print(f"[TODO] analyze: {m.run_tag}")
            ensured.append(m)
            continue

        err = ensure_analysis(root, m, params_path=params_path, force=bool(args.force))
        if err is None:
            ensured.append(m)
        else:
            failures.append(err)

    records: List[AnalysisRecord] = [load_record(root, m) for m in ensured]

    report_path = os.path.join(root, "ana_prev_n_class", args.report_name)
    if args.dry_run:
        print(f"[DRY] would write report: {report_path}")
    else:
        build_docx_report(records, out_path=report_path, params_path=params_path)
        print(f"[OK] Wrote report: {report_path}")

    if failures:
        print("\n[WARN] Some runs failed analysis:")
        for s in failures:
            print(" -", s)


if __name__ == "__main__":
    main()
