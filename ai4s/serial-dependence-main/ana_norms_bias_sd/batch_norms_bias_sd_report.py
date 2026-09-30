# ana_norms_bias_sd/batch_norms_bias_sd_report.py
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import subprocess
import sys
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
import yaml

# ---- optional tqdm (graceful fallback) ----
try:
    from tqdm import tqdm  # type: ignore
except Exception:
    class tqdm:  # minimal no-op fallback
        def __init__(self, iterable=None, total=None, desc=None, leave=False, **kwargs):
            self.iterable = iterable
            self.total = total
        def __iter__(self):
            return iter(self.iterable or [])
        def update(self, n=1):  # noqa
            pass
        def set_postfix(self, *args, **kwargs):  # noqa
            pass
        def close(self):  # noqa
            pass


SKIP_DIMS = ("none", "2d", "mixed")


def project_root_dir() -> str:
    here = os.path.abspath(os.path.dirname(__file__))  # .../ana_norms_bias_sd
    parent = os.path.abspath(os.path.join(here, ".."))
    if os.path.isdir(os.path.join(parent, "outputs")) or os.path.exists(os.path.join(parent, "config.yaml")):
        return parent
    if os.path.isdir(os.path.join(here, "outputs")) or os.path.exists(os.path.join(here, "config.yaml")):
        return here
    return parent


def safe_load_yaml(path: str) -> Dict[str, Any]:
    try:
        with open(path, "r", encoding="utf-8") as f:
            obj = yaml.safe_load(f) or {}
        if isinstance(obj, dict):
            return obj
    except Exception:
        pass
    return {}


def parse_n_from_run_tag(run_tag: str) -> Optional[int]:
    m = re.search(r"(?:^|_)n(\d+)(?:_|$)", run_tag)
    if not m:
        return None
    try:
        return int(m.group(1))
    except Exception:
        return None


def infer_model_from_run_tag(run_tag: str) -> Optional[str]:
    parts = run_tag.split("_")
    if len(parts) < 3:
        return None
    # often: 20260101_000000_<model>_<which>_<dim>_n400
    return parts[2]


def infer_which_from_run_tag(run_tag: str) -> Optional[str]:
    parts = run_tag.split("_")
    for p in parts:
        if p in ("exp1", "exp2", "both"):
            return p
    return None


def is_warriner_run(cfg: Dict[str, Any], run_tag: str) -> bool:
    # prefer explicit meta
    wsm = cfg.get("word_source_meta") or {}
    if isinstance(wsm, dict) and str(wsm.get("source_type", "")).lower() == "warriner_vad":
        return True

    excel_path = str(cfg.get("excel_path") or "").lower()
    if "warriner" in excel_path or "emotratings" in excel_path:
        return True

    # fallback: run_tag hints (rare)
    if "warriner" in run_tag.lower():
        return True

    return False


@dataclass
class RunMeta:
    run_tag: str
    run_dir: str
    model: str
    which: str
    dimension: str
    n_words: Optional[int]
    uses_warriner: bool


def extract_run_meta(run_dir: str) -> Optional[RunMeta]:
    run_tag = os.path.basename(os.path.normpath(run_dir))
    cfg_path = os.path.join(run_dir, "run_config.yaml")
    cfg = safe_load_yaml(cfg_path) if os.path.exists(cfg_path) else {}

    dimension = str(cfg.get("dimension") or "").strip().lower() or "unknown"
    if dimension in SKIP_DIMS:
        return None

    which = str(cfg.get("which") or "").strip().lower()
    if not which:
        which = infer_which_from_run_tag(run_tag) or "unknown"

    model = str(cfg.get("model") or "").strip()
    if not model:
        model = infer_model_from_run_tag(run_tag) or "(unknown)"

    n_words = None
    if cfg.get("n_words") is not None:
        try:
            n_words = int(cfg.get("n_words"))
        except Exception:
            n_words = None
    if n_words is None:
        n_words = parse_n_from_run_tag(run_tag)

    uses_w = is_warriner_run(cfg, run_tag)

    return RunMeta(
        run_tag=run_tag,
        run_dir=run_dir,
        model=model,
        which=which,
        dimension=dimension,
        n_words=n_words,
        uses_warriner=uses_w,
    )


def has_required_data(run_dir: str) -> bool:
    # analyzer can use exp2_sequential_raw.csv or combined_overview.csv
    p1 = os.path.join(run_dir, "exp2_sequential_raw.csv")
    p2 = os.path.join(run_dir, "combined_overview.csv")
    return os.path.exists(p1) or os.path.exists(p2)


def has_exp1_summary(run_dir: str) -> bool:
    p1 = os.path.join(run_dir, "exp1_isolated_summary.csv")
    p2 = os.path.join(run_dir, "combined_overview.csv")
    if os.path.exists(p1):
        return True
    if os.path.exists(p2):
        # combined_overview must contain mean
        try:
            df = pd.read_csv(p2, nrows=5)
            return "mean" in df.columns
        except Exception:
            return False
    return False


def analysis_dir(root: str, run_tag: str, mode: str) -> str:
    base = os.path.join(root, "ana_norms_bias_sd", "results")
    if mode == "exp1":
        return os.path.join(base, f"{run_tag}__exp1")
    return os.path.join(base, run_tag)


def analysis_paths(root: str, run_tag: str, mode: str) -> Dict[str, str]:
    d = analysis_dir(root, run_tag, mode)
    return {
        "dir": d,
        "summary_csv": os.path.join(d, "nback_summary.csv"),
        "fit_json": os.path.join(d, "fit_all_k.json"),
        "bias_json": os.path.join(d, "bias_fit.json"),
        "halfamp_plot": os.path.join(d, "halfamp_vs_k.png"),
        "e_vs_s_k1": os.path.join(d, "e_vs_s_k1.png"),
        "resid_scatter_k1": os.path.join(d, "e_resid_vs_delta_k1.png"),
        "resid_bin_k1": os.path.join(d, "e_resid_bin_vs_delta_k1.png"),
    }


def existing_perm_n(fit_json_path: str) -> Optional[int]:
    if not os.path.exists(fit_json_path):
        return None
    try:
        with open(fit_json_path, "r", encoding="utf-8") as f:
            obj = json.load(f)
        n_perm = obj.get("n_perm")
        if n_perm is None:
            return None
        return int(n_perm)
    except Exception:
        return None


def ensure_analysis(
    root: str,
    meta: RunMeta,
    config_path: str,
    min_n_perm: int,
    mode: str,
    force: bool = False,
) -> Tuple[bool, Optional[str]]:
    paths = analysis_paths(root, meta.run_tag, mode)

    ok_files = (
        os.path.exists(paths["summary_csv"])
        and os.path.exists(paths["halfamp_plot"])
        and os.path.exists(paths["e_vs_s_k1"])
        and os.path.exists(paths["resid_scatter_k1"])
        and os.path.exists(paths["resid_bin_k1"])
        and os.path.exists(paths["fit_json"])
        and os.path.exists(paths["bias_json"])
    )

    if ok_files and not force:
        prev_n = existing_perm_n(paths["fit_json"])
        if prev_n is not None and prev_n >= int(min_n_perm):
            return True, None
        # else re-run to upgrade permutations

    analyzer = os.path.join(root, "ana_norms_bias_sd", "analyze_norms_bias_sd.py")
    if not os.path.exists(analyzer):
        return False, f"Analyzer not found: {analyzer}"

    cmd = [
        sys.executable,
        analyzer,
        "--run-dir",
        meta.run_dir,
        "--config",
        config_path,
        "--n-perm",
        str(int(min_n_perm)),
    ]
    if mode == "exp1":
        cmd.append("--exp1")

    try:
        subprocess.run(cmd, cwd=root, check=True)
    except subprocess.CalledProcessError as e:
        return False, f"analyze_norms_bias_sd failed (run={meta.run_tag}, mode={mode}): {e}"

    ok_files2 = (
        os.path.exists(paths["summary_csv"])
        and os.path.exists(paths["halfamp_plot"])
        and os.path.exists(paths["e_vs_s_k1"])
        and os.path.exists(paths["resid_scatter_k1"])
        and os.path.exists(paths["resid_bin_k1"])
        and os.path.exists(paths["fit_json"])
        and os.path.exists(paths["bias_json"])
    )
    if ok_files2:
        return True, None
    return False, f"Analysis did not produce expected outputs for run={meta.run_tag} (mode={mode})"


@dataclass
class AnalysisRecord:
    meta: RunMeta
    mode: str
    stimulus_source: str
    summary_csv: Optional[str]
    halfamp_plot: Optional[str]
    e_vs_s_k1: Optional[str]
    resid_scatter_k1: Optional[str]
    resid_bin_k1: Optional[str]
    k1_half_amp: Optional[float]
    k1_p: Optional[float]
    k1_sigma: Optional[float]
    best_p: Optional[float]
    best_k: Optional[int]
    match_rate: Optional[float]
    note: Optional[str] = None


def load_record(root: str, meta: RunMeta, mode: str) -> AnalysisRecord:
    p = analysis_paths(root, meta.run_tag, mode)
    note = None

    stimulus_source = mode
    match_rate = None
    if os.path.exists(p["bias_json"]):
        try:
            with open(p["bias_json"], "r", encoding="utf-8") as f:
                bj = json.load(f)
            stimulus_source = str(bj.get("stimulus_source") or stimulus_source)
            mr = bj.get("match_rate_vs_exp2")
            if mr is not None:
                match_rate = float(mr)
        except Exception as e:
            note = f"Failed to parse bias_fit.json: {e}"

    k1_half_amp = None
    k1_p = None
    k1_sigma = None
    best_p = None
    best_k = None

    if os.path.exists(p["summary_csv"]):
        try:
            df = pd.read_csv(p["summary_csv"])
            if "k" in df.columns:
                df["k"] = df["k"].astype(int)

            dfk1 = df[df["k"] == 1] if "k" in df.columns else pd.DataFrame()
            if len(dfk1) == 1:
                row = dfk1.iloc[0]
                if pd.notna(row.get("half_amp")):
                    k1_half_amp = float(row.get("half_amp"))
                if pd.notna(row.get("p_half_amp")):
                    k1_p = float(row.get("p_half_amp"))
                if pd.notna(row.get("sigma")):
                    k1_sigma = float(row.get("sigma"))

            if "p_half_amp" in df.columns and "k" in df.columns:
                d2 = df.dropna(subset=["p_half_amp"]).copy()
                if not d2.empty:
                    idx = d2["p_half_amp"].astype(float).idxmin()
                    best_p = float(d2.loc[idx, "p_half_amp"])
                    best_k = int(d2.loc[idx, "k"])

        except Exception as e:
            note = (note + " ; " if note else "") + f"Failed to parse summary CSV: {e}"

    def _maybe(path: str) -> Optional[str]:
        return path if os.path.exists(path) else None

    return AnalysisRecord(
        meta=meta,
        mode=mode,
        stimulus_source=stimulus_source,
        summary_csv=_maybe(p["summary_csv"]),
        halfamp_plot=_maybe(p["halfamp_plot"]),
        e_vs_s_k1=_maybe(p["e_vs_s_k1"]),
        resid_scatter_k1=_maybe(p["resid_scatter_k1"]),
        resid_bin_k1=_maybe(p["resid_bin_k1"]),
        k1_half_amp=k1_half_amp,
        k1_p=k1_p,
        k1_sigma=k1_sigma,
        best_p=best_p,
        best_k=best_k,
        match_rate=match_rate,
        note=note,
    )


def build_docx_report(
    records: List[AnalysisRecord],
    out_path: str,
    params: Dict[str, Any],
) -> None:
    try:
        from docx import Document  # type: ignore
        from docx.shared import Inches  # type: ignore
    except Exception as e:
        raise RuntimeError(
            "python-docx is required to generate a Word report. Install via: pip install python-docx\n"
            f"Import error: {e}"
        )

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc = Document()

    doc.add_heading("Norms/Exp1 bias-removed serial dependence (DoG) report", level=0)
    doc.add_paragraph(f"Generated at: {_dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    doc.add_heading("Analysis parameters", level=1)
    doc.add_paragraph(f"k_values: {params.get('k_values')}")
    doc.add_paragraph(f"n_permutations: {params.get('n_permutations')}")
    doc.add_paragraph(f"random_seed: {params.get('random_seed')}")
    doc.add_paragraph(f"bias_method: {params.get('bias_method')}")

    doc.add_heading("Overview (k=1)", level=1)
    records2 = sorted(records, key=lambda r: (r.meta.run_tag, r.mode))

    if records2:
        tbl = doc.add_table(rows=1, cols=10)
        hdr = tbl.rows[0].cells
        hdr[0].text = "run_tag"
        hdr[1].text = "mode"
        hdr[2].text = "stimulus_source"
        hdr[3].text = "dimension"
        hdr[4].text = "model"
        hdr[5].text = "which"
        hdr[6].text = "n_words"
        hdr[7].text = "half_amp(k=1)"
        hdr[8].text = "p(k=1)"
        hdr[9].text = "match_rate"

        for r in records2:
            c = tbl.add_row().cells
            c[0].text = r.meta.run_tag
            c[1].text = r.mode
            c[2].text = r.stimulus_source
            c[3].text = r.meta.dimension
            c[4].text = r.meta.model
            c[5].text = r.meta.which
            c[6].text = str(r.meta.n_words) if r.meta.n_words is not None else ""
            c[7].text = f"{r.k1_half_amp:.6g}" if r.k1_half_amp is not None else ""
            c[8].text = f"{r.k1_p:.6g}" if r.k1_p is not None else ""
            c[9].text = f"{r.match_rate:.3g}" if r.match_rate is not None else ""
    else:
        doc.add_paragraph("(No eligible runs.)")

    # Detailed per run
    by_run: Dict[str, List[AnalysisRecord]] = {}
    for r in records2:
        by_run.setdefault(r.meta.run_tag, []).append(r)

    for run_tag in sorted(by_run.keys()):
        doc.add_heading(run_tag, level=1)
        for rec in by_run[run_tag]:
            doc.add_heading(f"Mode: {rec.mode}  (stimulus={rec.stimulus_source})", level=2)

            meta = rec.meta
            lines = [
                f"model: {meta.model}",
                f"which: {meta.which}",
                f"dimension: {meta.dimension}",
                f"uses_warriner_word_source: {meta.uses_warriner}",
            ]
            if meta.n_words is not None:
                lines.append(f"n_words: {meta.n_words}")
            if rec.match_rate is not None:
                lines.append(f"match_rate_vs_exp2: {rec.match_rate:.4g}")
            if rec.k1_half_amp is not None:
                lines.append(f"k=1 half_amp: {rec.k1_half_amp:.6g}")
            if rec.k1_sigma is not None:
                lines.append(f"k=1 sigma: {rec.k1_sigma:.6g}")
            if rec.k1_p is not None:
                lines.append(f"k=1 p (permutation): {rec.k1_p:.6g}")
            if rec.best_p is not None and rec.best_k is not None:
                lines.append(f"best p over ks: {rec.best_p:.6g} (k={rec.best_k})")
            if rec.note:
                lines.append(f"note: {rec.note}")

            doc.add_paragraph("\n".join(lines))

            # Figures
            if rec.halfamp_plot:
                doc.add_picture(rec.halfamp_plot, width=Inches(6.0))
            else:
                doc.add_paragraph("(missing halfamp_vs_k.png)")

            # 3 key plots
            tbl2 = doc.add_table(rows=1, cols=3)
            c0 = tbl2.cell(0, 0)
            c1 = tbl2.cell(0, 1)
            c2 = tbl2.cell(0, 2)

            if rec.e_vs_s_k1:
                p = c0.paragraphs[0]
                run = p.add_run()
                run.add_picture(rec.e_vs_s_k1, width=Inches(2.0))
            else:
                c0.text = "(missing e_vs_s_k1.png)"

            if rec.resid_scatter_k1:
                p = c1.paragraphs[0]
                run = p.add_run()
                run.add_picture(rec.resid_scatter_k1, width=Inches(2.0))
            else:
                c1.text = "(missing e_resid_vs_delta_k1.png)"

            if rec.resid_bin_k1:
                p = c2.paragraphs[0]
                run = p.add_run()
                run.add_picture(rec.resid_bin_k1, width=Inches(2.0))
            else:
                c2.text = "(missing e_resid_bin_vs_delta_k1.png)"

            doc.add_paragraph("")

    doc.save(out_path)


def decide_modes(meta: RunMeta) -> List[str]:
    # Using Warriner word source -> just norms
    if meta.uses_warriner:
        return ["norms"]

    dim = meta.dimension.lower()
    model = meta.model.lower()

    # your rule: non-warriner + gpt-4o + dim not in skip -> run both
    if ("gpt-4o" in model or model == "gpt4o") and dim not in SKIP_DIMS:
        return ["exp1", "norms"]

    # default: norms only
    return ["norms"]


def main() -> None:
    ap = argparse.ArgumentParser(description="Scan outputs, run norms/exp1 analysis, and generate a Word report.")
    ap.add_argument("--outputs-dir", type=str, default="outputs", help="Outputs directory under project root.")
    ap.add_argument("--report-dir", type=str, default="ana_norms_bias_sd", help="Directory (under project root) for report.")
    ap.add_argument("--config", type=str, default="ana_norms_bias_sd/params.yaml", help="params.yaml path (relative to root).")
    ap.add_argument("--report-name", type=str, default="norms_bias_sd_report.docx")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true", help="Force re-run analysis even if outputs exist.")
    args = ap.parse_args()

    root = project_root_dir()
    outputs_dir = os.path.join(root, args.outputs_dir)
    if not os.path.isdir(outputs_dir):
        raise FileNotFoundError(f"Outputs directory not found: {outputs_dir}")

    config_path = os.path.join(root, args.config)
    if not os.path.exists(config_path):
        raise FileNotFoundError(f"params.yaml not found: {config_path}")

    params = safe_load_yaml(config_path)
    min_n_perm = int(params.get("n_permutations") or 0)

    # scan runs
    run_metas: List[RunMeta] = []
    for name in sorted(os.listdir(outputs_dir)):
        run_dir = os.path.join(outputs_dir, name)
        if not os.path.isdir(run_dir):
            continue
        cfg_path = os.path.join(run_dir, "run_config.yaml")
        if not os.path.exists(cfg_path):
            continue
        meta = extract_run_meta(run_dir)
        if meta is None:
            continue
        if not has_required_data(run_dir):
            continue
        run_metas.append(meta)

    failures: List[str] = []
    records: List[AnalysisRecord] = []

    bar = tqdm(total=len(run_metas), desc="Analyze", leave=True)

    for meta in run_metas:
        modes = decide_modes(meta)

        # exp1 mode needs exp1 summary; otherwise skip exp1 cleanly
        if "exp1" in modes and not has_exp1_summary(meta.run_dir):
            failures.append(f"Missing exp1_isolated_summary.csv for exp1 analysis (run={meta.run_tag})")
            modes = [m for m in modes if m != "exp1"]

        # run each mode
        for mode in modes:
            if args.dry_run:
                bar.set_postfix({"run": meta.run_tag, "mode": mode, "dry": True})
                rec = load_record(root, meta, mode)
                records.append(rec)
                continue

            bar.set_postfix({"run": meta.run_tag, "mode": mode})
            ok, err = ensure_analysis(
                root=root,
                meta=meta,
                config_path=config_path,
                min_n_perm=min_n_perm,
                mode=mode,
                force=bool(args.force),
            )
            if ok:
                rec = load_record(root, meta, mode)
                records.append(rec)
            else:
                failures.append(err or f"unknown failure for run={meta.run_tag}, mode={mode}")

        bar.update(1)

    bar.close()

    report_path = os.path.join(root, args.report_dir, args.report_name)
    if args.dry_run:
        print(f"[DRY] would write report: {report_path}")
    else:
        build_docx_report(records, out_path=report_path, params=params)
        print(f"[OK] Wrote report: {report_path}")

    if failures:
        print("\n[WARN] Some runs failed analysis:")
        for s in failures:
            print(" -", s)


if __name__ == "__main__":
    main()
